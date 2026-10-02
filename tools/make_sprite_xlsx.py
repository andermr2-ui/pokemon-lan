# encoding: utf-8
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SRC = "tools/_sprite_inventory.csv"
OUT = "Inventario_Sprites_CharCustom.xlsx"

FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="2F5496")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=16, color="2F5496")
NOTE_FONT = Font(name=FONT_NAME, italic=True, size=10, color="595959")
BASE_FONT = Font(name=FONT_NAME, size=10)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
YELLOW = PatternFill("solid", fgColor="FFEB9C")
GREY = PatternFill("solid", fgColor="F2F2F2")

with open(SRC, newline="", encoding="utf-8") as f:
    reader = csv.reader(f)
    header = next(reader)
    rows = list(reader)

# En catalogo primero, despues alfabetico por personaje
rows.sort(key=lambda r: (r[2] != "Si", r[0]))

wb = Workbook()
ws = wb.active
ws.title = "Inventario"

# --- Titulo y leyenda ---
ws.merge_cells("A1:L1")
ws["A1"] = "Inventario de sprites — Character Customization"
ws["A1"].font = TITLE_FONT

ws.merge_cells("A2:L2")
ws["A2"] = (
    "Columnas de estado: si el sprite existe, la celda tiene la RUTA del archivo (relativa a la "
    "carpeta del juego). '(placeholder, copia del de caminar)' junto a la ruta significa que el "
    "archivo existe pero es identico al charset de caminar, sin animacion propia todavia. "
    "'(usa el de surf...)' en Buceo significa que no hace falta un archivo de buceo aparte: el motor "
    "usa el charset de surf automaticamente bajo el agua. 'Camina o esta quieto' compara las 4 poses "
    "de la fila de abajo del charset: si son todas iguales, el personaje no tiene animacion de "
    "caminar real todavia aunque el archivo exista. FALTA = no existe el archivo. Actualiza a mano "
    "la celda con la ruta cuando crees un sprite nuevo."
)
ws["A2"].font = NOTE_FONT
ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[2].height = 45

HEADER_ROW = 4
for col, title in enumerate(header, start=1):
    c = ws.cell(row=HEADER_ROW, column=col, value=title)
    c.font = HEADER_FONT
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORDER
ws.row_dimensions[HEADER_ROW].height = 30

STATUS_COLS = [3, 4, 6, 7, 8, 9, 10, 11]  # En catalogo, Charset caminar, run..fish, Back sprite
WALK_ANIM_COL = 5  # Camina o esta quieto (logica de color propia)
NOTAS_COL = 12

for i, row in enumerate(rows):
    r = HEADER_ROW + 1 + i
    for col, value in enumerate(row, start=1):
        c = ws.cell(row=r, column=col, value=value)
        c.font = BASE_FONT
        c.border = BORDER
        c.alignment = Alignment(vertical="center", wrap_text=(col in STATUS_COLS or col == WALK_ANIM_COL or col == NOTAS_COL))
        if col in STATUS_COLS:
            v = (value or "").strip()
            if v == "Si":
                c.fill = GREEN
            elif v == "FALTA":
                c.fill = RED
            elif v.startswith("N/A"):
                c.fill = GREY
            elif "(placeholder" in v:
                c.fill = YELLOW
            elif v:
                # tiene una ruta de archivo real (incluye "usa el de surf" para buceo)
                c.fill = GREEN
        elif col == WALK_ANIM_COL:
            v = (value or "").strip()
            if v == "Camina":
                c.fill = GREEN
            elif v.startswith("Quieto"):
                c.fill = YELLOW
            elif v.startswith("N/A"):
                c.fill = GREY
    if i % 2 == 1:
        for col in range(1, len(row) + 1):
            cell = ws.cell(row=r, column=col)
            if cell.fill.fgColor.rgb in (None, "00000000"):
                cell.fill = PatternFill("solid", fgColor="FAFAFA")

last_row = HEADER_ROW + len(rows)

# Las columnas de estado ahora llevan la ruta del archivo (texto libre) en
# vez de una palabra fija, asi que solo tiene sentido un desplegable en
# "En catalogo" (Si/No).
dv3 = DataValidation(type="list", formula1='"Si,No"', allow_blank=True)
ws.add_data_validation(dv3)
dv3.add(f"C{HEADER_ROW + 1}:C{last_row}")

widths = [22, 30, 11, 42, 26, 42, 42, 42, 48, 42, 42, 40]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.freeze_panes = f"A{HEADER_ROW + 1}"
ws.auto_filter.ref = f"A{HEADER_ROW}:L{last_row}"

wb.save(OUT)
print(f"guardado: {OUT} ({len(rows)} filas)")
